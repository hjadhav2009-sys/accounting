from __future__ import annotations

import hashlib
import json
import io
import zipfile
import re
from typing import Any
from uuid import UUID,uuid4


class MasterRepository:
    def __init__(self,connect):self.connect=connect
    @staticmethod
    def _records(cursor):
        names=[column.name if hasattr(column,"name") else column[0] for column in cursor.description]
        return [dict(zip(names,row)) for row in cursor.fetchall()]

    def snapshot(self,organization_id:UUID,company_id:UUID)->dict[str,Any]:
        connection=self.connect()
        try:
            with connection.cursor() as cursor:
                cursor.execute("SELECT id,name,tally_company_name,gstin,state,created_at FROM companies WHERE organization_id=%s ORDER BY name",(organization_id,));companies=self._records(cursor)
                cursor.execute("SELECT id,account_hint_token,bank_ledger,status FROM bank_accounts WHERE organization_id=%s AND company_id=%s ORDER BY bank_ledger",(organization_id,company_id));bank=self._records(cursor)
                for item in bank:item["account_hint"]="••••"+str(item.pop("account_hint_token"))[-4:]
                cursor.execute("SELECT id,platform,party_ledger,party_gstin,state FROM party_ledgers WHERE organization_id=%s AND company_id=%s ORDER BY platform",(organization_id,company_id));parties=self._records(cursor)
                cursor.execute("SELECT id,tax_type,ledger_name FROM gst_ledgers WHERE organization_id=%s AND company_id=%s ORDER BY tax_type",(organization_id,company_id));gst=self._records(cursor)
                cursor.execute("SELECT id,tool,platform,pattern,voucher_type,ledger,match_type,priority,enabled FROM ledger_mappings WHERE organization_id=%s AND company_id=%s ORDER BY tool,platform,priority DESC,pattern",(organization_id,company_id));mappings=self._records(cursor)
                cursor.execute("SELECT id,platform,document_type,tally_voucher_type,sign_mode FROM voucher_rules WHERE organization_id=%s AND company_id=%s ORDER BY platform,document_type",(organization_id,company_id));vouchers=self._records(cursor)
                return {"companies":companies,"bank_accounts":bank,"party_ledgers":parties,"gst_ledgers":gst,"ledger_mappings":mappings,"voucher_rules":vouchers}
        finally:connection.close()

    def _audit(self,cursor,organization_id,company_id,actor_id,action,entity_type,entity_id,summary):
        cursor.execute("""INSERT INTO audit_logs(id,organization_id,company_id,actor_id,action,entity_type,entity_id,
            new_reference,reason,source_context) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,'authorized masters update',%s::jsonb)""",
            (uuid4(),organization_id,company_id,actor_id,action,entity_type,str(entity_id),hashlib.sha256(json.dumps(summary,sort_keys=True).encode()).hexdigest(),json.dumps(summary)))

    def save_bank(self,organization_id,company_id,actor_id,account_hint,bank_ledger):
        hint="".join(character for character in account_hint if character.isalnum())
        if not 4<=len(hint)<=8:raise ValueError("account hint must contain 4 to 8 letters or digits")
        connection=self.connect();item_id=uuid4()
        try:
            with connection.cursor() as cursor:
                cursor.execute("""INSERT INTO bank_accounts(id,organization_id,company_id,account_hint_token,bank_ledger,status)
                    VALUES(%s,%s,%s,%s,%s,'ACTIVE') ON CONFLICT(company_id,account_hint_token,bank_ledger)
                    DO UPDATE SET status='ACTIVE' RETURNING id""",(item_id,organization_id,company_id,hint,bank_ledger.strip()));item_id=cursor.fetchone()[0]
                self._audit(cursor,organization_id,company_id,actor_id,"BANK_ACCOUNT_CHANGED","bank_account",item_id,{"account_hint":"masked","bank_ledger":bank_ledger.strip()})
            connection.commit();return item_id
        except Exception:connection.rollback();raise
        finally:connection.close()

    def save_party(self,organization_id,company_id,actor_id,platform,ledger,gstin="",state=""):
        connection=self.connect();item_id=uuid4()
        try:
            with connection.cursor() as cursor:
                cursor.execute("""INSERT INTO party_ledgers(id,organization_id,company_id,platform,party_ledger,party_gstin,state)
                    VALUES(%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(company_id,platform) DO UPDATE SET party_ledger=excluded.party_ledger,
                    party_gstin=excluded.party_gstin,state=excluded.state RETURNING id""",(item_id,organization_id,company_id,platform.strip().casefold(),ledger.strip(),gstin.strip() or None,state.strip() or None));item_id=cursor.fetchone()[0]
                self._audit(cursor,organization_id,company_id,actor_id,"PARTY_LEDGER_CHANGED","party_ledger",item_id,{"platform":platform,"ledger":ledger})
            connection.commit();return item_id
        except Exception:connection.rollback();raise
        finally:connection.close()

    def save_gst(self,organization_id,company_id,actor_id,tax_type,ledger):
        connection=self.connect();item_id=uuid4()
        try:
            with connection.cursor() as cursor:
                cursor.execute("""INSERT INTO gst_ledgers(id,organization_id,company_id,tax_type,ledger_name) VALUES(%s,%s,%s,%s,%s)
                    ON CONFLICT(company_id,tax_type) DO UPDATE SET ledger_name=excluded.ledger_name RETURNING id""",(item_id,organization_id,company_id,tax_type,ledger.strip()));item_id=cursor.fetchone()[0]
                self._audit(cursor,organization_id,company_id,actor_id,"GST_LEDGER_CHANGED","gst_ledger",item_id,{"tax_type":tax_type,"ledger":ledger})
            connection.commit();return item_id
        except Exception:connection.rollback();raise
        finally:connection.close()

    def save_voucher(self,organization_id,company_id,actor_id,platform,document_type,voucher_type,sign_mode):
        connection=self.connect();item_id=uuid4()
        try:
            with connection.cursor() as cursor:
                cursor.execute("""INSERT INTO voucher_rules(id,organization_id,company_id,platform,document_type,tally_voucher_type,sign_mode)
                    VALUES(%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(company_id,platform,document_type) DO UPDATE SET
                    tally_voucher_type=excluded.tally_voucher_type,sign_mode=excluded.sign_mode RETURNING id""",(item_id,organization_id,company_id,platform.strip().casefold(),document_type.strip(),voucher_type.strip(),sign_mode));item_id=cursor.fetchone()[0]
                self._audit(cursor,organization_id,company_id,actor_id,"VOUCHER_RULE_CHANGED","voucher_rule",item_id,{"platform":platform,"document_type":document_type,"voucher_type":voucher_type,"sign_mode":sign_mode})
            connection.commit();return item_id
        except Exception:connection.rollback();raise
        finally:connection.close()

    def mapping_workbook(self,organization_id:UUID,company_id:UUID)->bytes:
        from openpyxl import Workbook
        snapshot=self.snapshot(organization_id,company_id);workbook=Workbook();sheet=workbook.active;sheet.title="Ledger Mappings"
        columns=["tool","platform","pattern","voucher_type","ledger","match_type","priority","enabled"];sheet.append(columns)
        for item in snapshot["ledger_mappings"]:sheet.append([item.get(column,"") for column in columns])
        sheet.freeze_panes="A2";widths=[16,18,45,20,35,18,10,10]
        for index,width in enumerate(widths,start=1):sheet.column_dimensions[chr(64+index)].width=width
        stream=io.BytesIO();workbook.save(stream);return stream.getvalue()

    def preview_mapping_workbook(self,organization_id:UUID,company_id:UUID,actor_id:UUID,content:bytes)->dict[str,Any]:
        if not content or len(content)>5*1024*1024:raise ValueError("mapping workbook must be between 1 byte and 5 MB")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist())>200 or sum(item.file_size for item in archive.infolist())>25*1024*1024:raise ValueError("mapping workbook expands beyond safe limits")
        except zipfile.BadZipFile as exc:raise ValueError("mapping workbook is not a valid XLSX file") from exc
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        try:sheet=workbook["Ledger Mappings"] if "Ledger Mappings" in workbook.sheetnames else workbook.active;values=sheet.iter_rows(values_only=True);headers=[str(value or "").strip() for value in next(values,())]
        finally:pass
        required=["tool","platform","pattern","voucher_type","ledger","match_type","priority","enabled"]
        if headers!=required:workbook.close();raise ValueError("mapping workbook columns do not match the exported template")
        current=self.snapshot(organization_id,company_id)["ledger_mappings"];existing={(item["tool"],item["platform"],item["pattern"],item["voucher_type"]):item for item in current}
        rows=[];seen={};summary={"inserted":0,"updated":0,"unchanged":0,"invalid":0,"conflicts":0}
        for number,values_row in enumerate(values,start=2):
            if number>5001:workbook.close();raise ValueError("mapping workbook contains more than 5000 rows")
            row=dict(zip(required,values_row));row={key:("" if value is None else value) for key,value in row.items()};row["tool"]=str(row["tool"]).strip().casefold();row["platform"]=str(row["platform"]).strip().casefold();row["pattern"]=str(row["pattern"]).strip();row["voucher_type"]=str(row["voucher_type"]).strip();row["ledger"]=str(row["ledger"]).strip();row["match_type"]=str(row["match_type"]).strip();row["priority"]=int(row["priority"] or 0);row["enabled"]=str(row["enabled"]).casefold() not in {"false","0","no","off"}
            errors=[]
            if row["tool"] not in {"marketplace","bank"}:errors.append("tool")
            if not row["pattern"]:errors.append("blank pattern")
            if not row["ledger"]:errors.append("blank ledger")
            if row["match_type"] not in {"contains","smart_contains","equals","starts_with","regex"}:errors.append("match type")
            if row["match_type"]=="regex":
                try:re.compile(row["pattern"])
                except re.error:errors.append("invalid regex")
            key=(row["tool"],row["platform"],row["pattern"],row["voucher_type"]);prior=seen.get(key)
            if prior and prior!=row:errors.append("conflicting duplicate");summary["conflicts"]+=1;status="conflict"
            elif errors:summary["invalid"]+=1;status="invalid"
            elif key not in existing:summary["inserted"]+=1;status="inserted"
            else:
                old=existing[key];changed=any(str(old.get(field,"")).casefold()!=str(row[field]).casefold() for field in ("ledger","match_type","priority","enabled"))
                status="updated" if changed else "unchanged";summary[status]+=1
            seen[key]=row;rows.append({**row,"row_number":number,"status":status,"errors":errors})
        workbook.close();preview_id=uuid4();digest=hashlib.sha256(content).hexdigest();connection=self.connect()
        try:
            with connection.cursor() as cursor:cursor.execute("""INSERT INTO mapping_import_previews(id,organization_id,company_id,workbook_sha256,rows_json,summary,created_by,expires_at)
                VALUES(%s,%s,%s,%s,%s::jsonb,%s::jsonb,%s,now()+interval '30 minutes')""",(preview_id,organization_id,company_id,digest,json.dumps(rows),json.dumps(summary),actor_id))
            connection.commit()
        finally:connection.close()
        return {"preview_id":preview_id,"workbook_sha256":digest,"summary":summary,"rows":rows,"apply_allowed":summary["invalid"]==0 and summary["conflicts"]==0}

    def apply_mapping_preview(self,organization_id:UUID,company_id:UUID,actor_id:UUID,preview_id:UUID)->dict[str,Any]:
        connection=self.connect()
        try:
            with connection.cursor() as cursor:
                cursor.execute("SELECT rows_json,summary FROM mapping_import_previews WHERE id=%s AND organization_id=%s AND company_id=%s AND applied_at IS NULL AND expires_at>now() FOR UPDATE",(preview_id,organization_id,company_id));preview=cursor.fetchone()
                if not preview:raise ValueError("mapping preview is missing, expired, or already applied")
                rows,summary=preview
                if summary.get("invalid") or summary.get("conflicts"):raise ValueError("mapping preview contains invalid or conflicting rows")
                cursor.execute("SELECT tool,platform,pattern,voucher_type,ledger,match_type,priority,enabled FROM ledger_mappings WHERE organization_id=%s AND company_id=%s ORDER BY id",(organization_id,company_id));columns=[column.name for column in cursor.description];snapshot=[dict(zip(columns,row)) for row in cursor.fetchall()];serialized=json.dumps(snapshot,sort_keys=True);backup_id=uuid4()
                cursor.execute("INSERT INTO mapping_import_backups(id,organization_id,company_id,preview_id,snapshot,snapshot_sha256,created_by) VALUES(%s,%s,%s,%s,%s::jsonb,%s,%s)",(backup_id,organization_id,company_id,preview_id,serialized,hashlib.sha256(serialized.encode()).hexdigest(),actor_id))
                for row in rows:
                    if row["status"]=="unchanged":continue
                    cursor.execute("""INSERT INTO ledger_mappings(id,organization_id,company_id,tool,platform,pattern,voucher_type,ledger,match_type,priority,enabled)
                        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(company_id,tool,platform,pattern,voucher_type)
                        DO UPDATE SET ledger=excluded.ledger,match_type=excluded.match_type,priority=excluded.priority,enabled=excluded.enabled""",
                        (uuid4(),organization_id,company_id,row["tool"],row["platform"],row["pattern"],row["voucher_type"],row["ledger"],row["match_type"],row["priority"],row["enabled"]))
                cursor.execute("UPDATE mapping_import_previews SET applied_at=now() WHERE id=%s",(preview_id,));self._audit(cursor,organization_id,company_id,actor_id,"MAPPING_IMPORT_APPLIED","mapping_import",preview_id,{"summary":summary,"backup_id":str(backup_id)})
            connection.commit();return {"status":"APPLIED","summary":summary,"backup_id":backup_id}
        except Exception:connection.rollback();raise
        finally:connection.close()
